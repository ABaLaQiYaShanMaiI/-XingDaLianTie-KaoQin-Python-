# -*- coding: utf-8 -*-
"""GUI 冒烟 + 全流程测试: 建界面 -> 添加照片 -> 同步识别 -> 同步生成 -> 校验。"""
import os
import queue
import sys

import openpyxl

import _test_roster as fixture
import app as app_mod

BASE = os.path.dirname(os.path.abspath(__file__))
PHOTO = os.path.join(BASE, '参考文件', '考核明细图片.jpg')
OUT_DIR = os.path.join(BASE, '_out')


def drain(q):
    msgs = []
    try:
        while True:
            msgs.append(q.get_nowait())
    except queue.Empty:
        pass
    return msgs


def main():
    cfg_path = os.path.join(BASE, 'paofen_config.json')
    saved_cfg = None                       # 测试前快照(结束后还原用户配置)
    if os.path.isfile(cfg_path):
        with open(cfg_path, 'r', encoding='utf-8') as f:
            saved_cfg = f.read()
    root = app_mod.TkinterDnD.Tk() if app_mod.TkinterDnD else None
    if root is None:
        import tkinter as tk
        root = tk.Tk()
    gui = app_mod.PaofenApp(root)
    gui.roster_files.clear()        # 测试从干净状态开始(名单用模板内默认)
    gui._persist_config()
    root.update()

    # 默认输出目录应为桌面(取不到时退回程序目录)
    cfg_before = dict(gui.cfg)
    gui.cfg.pop('out_dir', None)               # 模拟首次运行(无记忆)
    gui._load_initial_config()
    assert gui.var_out.get() == app_mod.default_out_dir(), \
        f'默认输出目录异常: {gui.var_out.get()}'
    print('默认输出目录:', gui.var_out.get())
    gui.cfg.update(cfg_before)
    gui._load_initial_config()

    # 配置
    gui.var_template.set(os.path.join(BASE, '参考文件',
                                      '煤库作业区跑分绩效汇总X月（表样）.xlsx'))
    gui.var_out.set(OUT_DIR)
    gui.var_month.set('')
    gui.var_pool.set('')
    gui._add_photos([PHOTO])
    assert len(gui.photos) == 1, '照片添加失败'

    # 同步识别(直接调工作函数, 等价于后台线程执行体)
    gui._identify_worker()
    msgs = drain(gui.q)
    res = gui.results[PHOTO]
    print('识别: month=%s shift=%s 条目=%d 警告=%d'
          % (res['month'], res['shift'], len(res['entries']), len(res['warnings'])))
    assert res['month'] == 7 and res['shift'] == '甲班'
    assert len(res['entries']) == 15, f'条目数异常: {len(res["entries"])}'
    root.update()
    gui._refresh_photo_tree()
    gui._refresh_preview()
    assert len(gui.tv_entries.get_children()) == 15, '预览表行数异常'

    # 同步生成(屏蔽弹窗避免测试阻塞), 用 _poll 消化队列并更新界面
    app_mod.messagebox.showinfo = lambda *a, **k: None
    app_mod.messagebox.showerror = lambda *a, **k: None
    gui._generate_worker(gui.var_template.get(), [], 80, None, None)
    gui._poll()
    root.update()
    out_path = gui.out_path
    print('生成:', out_path)
    assert out_path and os.path.isfile(out_path), '输出文件不存在'
    assert '7月' in os.path.basename(out_path), '文件名月份异常'

    assert str(gui.btn_open_file['state']) == 'normal'
    assert str(gui.btn_open_dir['state']) == 'normal'

    # ---- 场景2: 仅「报名表(名单变动 20 人) + 手工录入明细」生成 ----
    gui._clear_photos()                          # 清空照片, 只用文本明细
    names20 = fixture.ROSTER_JULY + ['何建军', '吴建兵']
    roster_path = os.path.join(OUT_DIR, '_gui_roster20.xlsx')
    wb0 = openpyxl.Workbook()
    ws0 = wb0.active
    ws0.title = '甲班名单'
    ws0.append(['序号', '姓名', '岗位'])
    for i, n in enumerate(names20, 1):
        ws0.append([i, n, '皮带工'])
    wb0.save(roster_path)
    wb0.close()

    gui.roster_files['甲班'] = roster_path       # 等价于界面「导入名单…」
    gui._persist_config()
    gui._refresh_roster_tree()
    row = [gui.tv_roster.item(i, 'values') for i in gui.tv_roster.get_children()]
    jia = [v for v in row if v[0] == '甲班'][0]
    print('名单表甲班行:', jia)
    assert str(jia[1]).endswith('_gui_roster20.xlsx'), jia
    assert int(jia[2]) == 20, jia
    assert app_mod.load_config().get('rosters', {}).get('甲班') == roster_path, \
        '选择未被记忆为默认'

    gui.var_manual_shift.set('甲班')             # 等价于界面「解析文本明细」
    gui.txt_manual.delete('1.0', 'end')
    gui.txt_manual.insert('1.0', fixture.DETAIL_TEXT)
    gui._parse_manual_text()
    root.update()
    assert len(gui.manual_results) == 1
    res_m = gui.manual_results[0]
    assert len(res_m['entries']) == 15, len(res_m['entries'])
    assert not res_m['warnings'], res_m['warnings']
    assert next(e for e in res_m['entries']
                if e['name_raw'].startswith('陈金荣'))['name'] == '陈金荣'
    assert len(gui.tv_entries.get_children()) == 15, '预览未显示手工明细'

    rosters, override, _notes = gui._effective_rosters()
    assert override.get('甲班') == names20, '报名表未生效'
    gui._generate_worker(gui.var_template.get(), [], 80, 10800, 7, rosters, override)
    gui._poll()
    root.update()
    out2 = gui.out_path
    print('生成(报名表+手工明细):', out2)
    assert out2 and out2 != out_path and os.path.isfile(out2)
    wb = openpyxl.load_workbook(out2)
    ws = wb['甲班']
    got_names = [str(ws.cell(row=r, column=2).value or '').strip()
                 for r in range(3, 23)]
    assert got_names == names20, f'考勤表名单未按报名表重排: {got_names}'
    assert ws['AH23'].value == '=SUM(AH3:AH22)', ws['AH23'].value
    assert ws['AI23'].value == 10800 and ws['AJ23'].value == '=AI23/AH23'
    assert ws.cell(row=3, column=35).value == '=ROUND(AH3*AJ$23,2)'
    assert ws.cell(row=10, column=20).value == 3          # 陈金荣 7.18(T列)
    assert ws.cell(row=13, column=25).value == -3         # 田忠山 7.23(Y列)
    wb.close()
    print('报名表重排/公式/日格 校验通过')

    gui.roster_files.pop('甲班')                 # 恢复模板名单(清理)
    gui._persist_config()
    root.destroy()
    if saved_cfg is None:                        # 测试前无配置 -> 删掉测试产生的
        if os.path.isfile(cfg_path):
            os.remove(cfg_path)
    else:                                        # 还原用户原配置(模板/输出目录/名单记忆)
        with open(cfg_path, 'w', encoding='utf-8') as f:
            f.write(saved_cfg)
    print('\nGUI 全流程测试通过')


if __name__ == '__main__':
    sys.exit(main())
