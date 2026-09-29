# -*- coding: utf-8 -*-
"""
app.py - 跑分绩效汇总自动生成工具(入口 + GUI)

界面/流程参考《安全生产责任制履职清单考评表自动生成工具》:
  tkinter + tkinterdnd2 图形界面; 照片拖拽/点选添加; 后台线程 OCR 不卡界面;
  识别结果预览(可人工核对); 一键生成 xlsx; 完成后一键打开文件/文件夹;
  输出与模板版式完全一致(公式实时计算), 模板只读、生成用副本。

核心逻辑:
  detail_core.py  考核明细照片 OCR 解析(标题/条目/姓名匹配)
  excel_core.py   模板名单读取与按明细填表(日格/合计/金额公式/明细存档)
"""
import json
import logging
import os
import queue
import sys
import threading
import time
import traceback
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import detail_core
import excel_core

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
except ImportError:  # 未装 tkinterdnd2 时退化为普通 tkinter(无拖拽)
    DND_FILES = None
    TkinterDnD = None

HERE = os.path.dirname(os.path.abspath(__file__))
if getattr(sys, 'frozen', False):      # PyInstaller 单文件运行时, 以 exe 所在目录为准
    HERE = os.path.dirname(os.path.abspath(sys.executable))
CONFIG_PATH = os.path.join(HERE, 'paofen_config.json')
LOG_PATH = os.path.join(HERE, 'paofen.log')
IMAGE_EXTS = ('.jpg', '.jpeg', '.png', '.bmp', '.webp')

logger = logging.getLogger('paofen')
logger.setLevel(logging.INFO)
_fh = logging.FileHandler(LOG_PATH, encoding='utf-8')
_fh.setFormatter(logging.Formatter('%(asctime)s %(levelname)s %(message)s'))
logger.addHandler(_fh)
logging.getLogger('rapidocr_onnxruntime').setLevel(logging.WARNING)

DEFAULT_CONFIG = {
    'template': '',
    'out_dir': '',          # 空=用默认(桌面); 记录上次选择
    'pool_amount': '',
    'base_score': excel_core.BASE_SCORE_DEFAULT,
    'rosters': {},          # {'甲班': '报名表.xlsx', ...} 上次选择的名单(默认记忆)
}


def _desktop_candidates():
    """桌面目录候选(Windows 优先读注册表, 兼容 OneDrive 重定向的桌面)。"""
    cands = []
    try:                                        # 1) 用户 Shell 文件夹(含重定向)
        import winreg
        with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                r'Software\Microsoft\Windows\CurrentVersion'
                r'\Explorer\User Shell Folders') as key:
            cands.append(os.path.expandvars(str(winreg.QueryValueEx(key, 'Desktop')[0])))
    except (ImportError, OSError, ValueError):
        pass
    try:                                        # 2) Windows API(CSIDL_DESKTOPDIRECTORY)
        import ctypes
        buf = ctypes.create_unicode_buffer(260)
        if ctypes.windll.shell32.SHGetFolderPathW(None, 0x0010, None, 0, buf) == 0:
            cands.append(buf.value)
    except (ImportError, AttributeError, OSError):
        pass
    cands.append(os.path.join(os.path.expanduser('~'), 'Desktop'))   # 3) 兜底
    return cands


def default_out_dir():
    """默认输出目录: 桌面(取不到时退回程序目录)。"""
    for p in _desktop_candidates():
        if p and os.path.isdir(p):
            return p
    return HERE


def load_config():
    cfg = dict(DEFAULT_CONFIG)
    try:
        with open(CONFIG_PATH, 'r', encoding='utf-8') as f:
            cfg.update(json.load(f))
    except (OSError, ValueError):
        pass
    return cfg


def save_config(cfg):
    try:
        with open(CONFIG_PATH, 'w', encoding='utf-8') as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
    except OSError as e:
        logger.warning('配置保存失败: %s', e)


class PaofenApp:
    def __init__(self, root):
        self.root = root
        root.title('跑分绩效汇总自动生成工具')
        root.geometry('1060x860')
        root.minsize(900, 640)
        self.cfg = load_config()
        self.photos = []      # [{'path': str, 'override': ''|'甲班'...}]
        self.results = {}     # path -> parse_photo 结果
        self.roster_files = {k: v for k, v in (self.cfg.get('rosters') or {}).items() if v}
        self.manual_text = {}         # {班次: 手工录入的明细文本}
        self.manual_results = []      # [parse 结果(dict), ...] 与照片结果同构
        self._tpl_rosters = None      # 模板内名单缓存(显示人数用)
        self.q = queue.Queue()
        self.busy = False
        self.out_path = None
        self._build_ui()
        self._load_initial_config()
        self.root.after(120, self._poll)

    # ---------------- 界面 ----------------
    def _build_ui(self):
        pad = dict(padx=8, pady=4)
        # ---- 顶部: 基本信息 ----
        top = ttk.LabelFrame(self.root, text=' 基本信息 ')
        top.pack(fill='x', **pad)
        ttk.Label(top, text='模板文件:').grid(row=0, column=0, sticky='e', **pad)
        self.var_template = tk.StringVar()
        ttk.Entry(top, textvariable=self.var_template).grid(
            row=0, column=1, sticky='we', **pad)
        ttk.Button(top, text='选择…', width=8,
                   command=self._pick_template).grid(row=0, column=2, **pad)
        ttk.Label(top, text='输出目录:').grid(row=1, column=0, sticky='e', **pad)
        self.var_out = tk.StringVar()
        ttk.Entry(top, textvariable=self.var_out).grid(
            row=1, column=1, sticky='we', **pad)
        ttk.Button(top, text='选择…', width=8,
                   command=self._pick_outdir).grid(row=1, column=2, **pad)
        ttk.Button(top, text='输出到桌面', width=10,
                   command=self._use_desktop_out).grid(row=1, column=5,
                                                       columnspan=2, sticky='w', **pad)
        ttk.Label(top, text='月份(空=自动):').grid(row=0, column=3, sticky='e', **pad)
        self.var_month = tk.StringVar()
        ttk.Entry(top, textvariable=self.var_month, width=8).grid(row=0, column=4, **pad)
        ttk.Label(top, text='基础分:').grid(row=1, column=3, sticky='e', **pad)
        self.var_base = tk.StringVar()
        ttk.Entry(top, textvariable=self.var_base, width=8).grid(row=1, column=4, **pad)
        ttk.Label(top, text='绩效总额(空=模板值):').grid(row=0, column=5, sticky='e', **pad)
        self.var_pool = tk.StringVar()
        ttk.Entry(top, textvariable=self.var_pool, width=12).grid(row=0, column=6, **pad)
        top.columnconfigure(1, weight=1)

        # ---- 名单(报名表) / 手工录入明细 ----
        nb = ttk.Notebook(self.root)
        nb.pack(fill='x', **pad)

        tab_r = ttk.Frame(nb)                      # 名单(报名表)
        nb.add(tab_r, text=' 名单(报名表) ')
        cols = ('shift', 'source', 'count', 'status')
        self.tv_roster = ttk.Treeview(tab_r, columns=cols, show='headings', height=4)
        for c, w, t in (('shift', 50, '班次'), ('source', 200, '名单来源'),
                        ('count', 50, '人数'), ('status', 290, '说明')):
            self.tv_roster.heading(c, text=t)
            self.tv_roster.column(c, width=w, anchor='w')
        self.tv_roster.pack(side='left', fill='both', expand=True, padx=(6, 0), pady=6)
        rb = ttk.Frame(tab_r)
        rb.pack(side='left', fill='y', padx=6, pady=6)
        ttk.Button(rb, text='导入名单…', command=self._pick_roster).pack(fill='x', pady=2)
        ttk.Button(rb, text='用模板名单', command=self._use_template_roster
                   ).pack(fill='x', pady=2)
        ttk.Label(tab_r, foreground='#666', justify='left', text=(
            '选中一行后点左侧按钮;\n导入即生效并记忆为默认;\n考勤表人名按报名表重排')).pack(
                side='left', padx=6)

        tab_m = ttk.Frame(nb)                      # 手工录入/粘贴明细
        nb.add(tab_m, text=' 手工录入/粘贴明细 ')
        mr = ttk.Frame(tab_m)
        mr.pack(fill='x', padx=6, pady=(6, 0))
        ttk.Label(mr, text='班次:').pack(side='left')
        self.var_manual_shift = tk.StringVar(value=detail_core.SHIFT_NAMES[0])
        ttk.Combobox(mr, textvariable=self.var_manual_shift, width=6, state='readonly',
                     values=detail_core.SHIFT_NAMES).pack(side='left', padx=4)
        ttk.Button(mr, text='解析文本明细', command=self._parse_manual_text
                   ).pack(side='left', padx=4)
        ttk.Label(mr, foreground='#666', text=(
            '一行一条, 如「1、7.2叶么菊发现现场安全隐患+3」; 序号可有可无, '
            '可与照片混用')).pack(side='left', padx=6)
        self.txt_manual = tk.Text(tab_m, height=5, wrap='none')
        self.txt_manual.pack(fill='both', expand=True, padx=6, pady=6)

        # ---- 中部: 明细照片 ----
        mid = ttk.LabelFrame(self.root, text=' 考核明细照片(每班一张, 可拖入) ')
        mid.pack(fill='both', expand=True, **pad)
        cols = ('file', 'shift', 'count', 'status')
        self.tv_photos = ttk.Treeview(mid, columns=cols, show='headings', height=4)
        for c, w, t in (('file', 380, '照片文件'), ('shift', 90, '班次'),
                        ('count', 70, '条目数'), ('status', 300, '状态')):
            self.tv_photos.heading(c, text=t)
            self.tv_photos.column(c, width=w, anchor='w')
        self.tv_photos.pack(side='left', fill='both', expand=True, padx=(6, 0), pady=6)
        vsb = ttk.Scrollbar(mid, orient='vertical', command=self.tv_photos.yview)
        vsb.pack(side='left', fill='y', pady=6)
        self.tv_photos.configure(yscrollcommand=vsb.set)
        if DND_FILES:
            self.tv_photos.drop_target_register(DND_FILES)
            self.tv_photos.dnd_bind('<<Drop>>', self._on_drop)
        btns = ttk.Frame(mid)
        btns.pack(side='left', fill='y', padx=6, pady=6)
        ttk.Button(btns, text='添加照片…', command=self._pick_photos).pack(fill='x', pady=2)
        ttk.Button(btns, text='移除选中', command=self._remove_photo).pack(fill='x', pady=2)
        ttk.Button(btns, text='清空全部', command=self._clear_photos).pack(fill='x', pady=2)
        ttk.Separator(btns, orient='horizontal').pack(fill='x', pady=6)
        self.var_override = tk.StringVar(value='自动')
        ov = ttk.Combobox(btns, textvariable=self.var_override, width=8,
                          state='readonly', values=('自动',) + detail_core.SHIFT_NAMES)
        ov.pack(fill='x', pady=2)
        ttk.Button(btns, text='指定选中\n照片班次', command=self._apply_override
                   ).pack(fill='x', pady=2)
        ttk.Button(btns, text='识别照片', command=self.start_identify
                   ).pack(fill='x', pady=(10, 2))

        # ---- 中下: 识别结果预览 ----
        prev = ttk.LabelFrame(self.root, text=' 识别结果预览(可人工核对) ')
        prev.pack(fill='both', expand=True, **pad)
        cols = ('photo', 'shift', 'seq', 'date', 'name', 'match', 'reason', 'score')
        self.tv_entries = ttk.Treeview(prev, columns=cols, show='headings', height=6)
        for c, w, t in (('photo', 70, '照片'), ('shift', 50, '班次'),
                        ('seq', 40, '序号'), ('date', 55, '日期'),
                        ('name', 90, '姓名(识别)'), ('match', 90, '匹配名单'),
                        ('reason', 260, '事由'), ('score', 55, '分值')):
            self.tv_entries.heading(c, text=t)
            self.tv_entries.column(c, width=w, anchor='w')
        self.tv_entries.pack(fill='both', expand=True, padx=6, pady=(6, 2))
        self.txt_warn = tk.Text(prev, height=4, wrap='word', fg='#a33')
        self.txt_warn.pack(fill='x', padx=6, pady=(0, 6))
        self.txt_warn.insert('1.0', '警告/提示会显示在这里')
        self.txt_warn.config(state='disabled')

        # ---- 底部: 进度与操作 ----
        bottom = ttk.Frame(self.root)
        bottom.pack(fill='x', **pad)
        self.progress = ttk.Progressbar(bottom, mode='determinate', length=260)
        self.progress.pack(side='left', padx=(8, 4))
        self.var_status = tk.StringVar(value='请添加考核明细照片(支持拖拽)')
        ttk.Label(bottom, textvariable=self.var_status, width=52, anchor='w'
                  ).pack(side='left', padx=4)
        self.btn_open_file = ttk.Button(bottom, text='打开文件', state='disabled',
                                        command=self._open_file)
        self.btn_open_dir = ttk.Button(bottom, text='打开所在文件夹', state='disabled',
                                       command=self._open_folder)
        self.btn_generate = ttk.Button(bottom, text=' 一键生成 ', command=self.start_generate)
        self.btn_open_file.pack(side='right', padx=4, pady=6)
        self.btn_open_dir.pack(side='right', padx=4, pady=6)
        self.btn_generate.pack(side='right', padx=(4, 10), pady=6)

    # ---------------- 配置/文件选择 ----------------
    def _load_initial_config(self):
        self.var_template.set(self.cfg.get('template') or '')
        if not self.var_template.get():
            t = excel_core.find_template(HERE)
            if t:
                self.var_template.set(t)
        saved_out = str(self.cfg.get('out_dir') or '').strip()
        if saved_out and os.path.isdir(saved_out):
            self.var_out.set(saved_out)          # 记忆上次选择
        else:
            if saved_out:                        # 换机器/目录被删 -> 回到默认
                logger.warning('输出目录不存在, 已改用默认目录: %s', saved_out)
            self.var_out.set(default_out_dir())
        self.var_pool.set(str(self.cfg.get('pool_amount') or ''))
        self.var_base.set(str(self.cfg.get('base_score') or excel_core.BASE_SCORE_DEFAULT))
        self._refresh_roster_tree()      # 上次选择的报名表(默认)带出来

    def _persist_config(self):
        self.cfg.update({
            'template': self.var_template.get().strip(),
            'out_dir': self.var_out.get().strip(),
            'pool_amount': self.var_pool.get().strip(),
            'base_score': self.var_base.get().strip(),
            'rosters': {k: v for k, v in (self.roster_files or {}).items() if v},
        })
        save_config(self.cfg)

    def _pick_template(self):
        p = filedialog.askopenfilename(title='选择跑分绩效汇总模板',
                                       filetypes=[('Excel 模板', '*.xlsx')])
        if p:
            self.var_template.set(p)
            self._tpl_rosters = None
            self._persist_config()
            self._refresh_roster_tree()

    def _pick_outdir(self):
        p = filedialog.askdirectory(title='选择输出目录')
        if p:
            self.var_out.set(p)
            self._persist_config()

    def _use_desktop_out(self):
        d = default_out_dir()
        self.var_out.set(d)
        self._persist_config()
        self._set_status(f'输出目录已设为桌面: {d}')

    # ---------------- 名单(报名表) / 手工录入明细 ----------------
    def _tpl_roster_map(self):
        """模板内名单(按模板路径缓存) -> {班次: [姓名]}。"""
        template = self.var_template.get().strip()
        if self._tpl_rosters is None or getattr(self, '_tpl_path', None) != template:
            try:
                self._tpl_rosters = (excel_core.load_rosters(template)
                                     if os.path.isfile(template) else {})
            except Exception as e:
                logger.warning('读取模板名单失败: %s', e)
                self._tpl_rosters = {}
            self._tpl_path = template
        return self._tpl_rosters

    def _effective_rosters(self):
        """生效名单(报名表覆盖模板名单)。

        返回 (rosters, override, notes):
          rosters  {班次: [姓名]} —— 姓名匹配用
          override {班次: [姓名]} —— 需按报名表重排考勤表名单的班次
          notes    [str]          —— 预览区提示
        """
        rosters = dict(self._tpl_roster_map())
        override, notes = {}, []
        for shift, path in sorted((self.roster_files or {}).items()):
            if not path:
                continue
            if not os.path.isfile(path):
                notes.append(f'{shift} 报名表不存在({path}), 已改用模板内名单')
                continue
            try:
                names, sheet = excel_core.load_roster_file(path, shift)
            except Exception as e:
                notes.append(f'{shift} 报名表读取失败({e}), 已改用模板内名单')
                continue
            if not names:
                notes.append(f'{shift} 报名表「{os.path.basename(path)}」未读到姓名, '
                             f'已改用模板内名单')
                continue
            rosters[shift] = names
            override[shift] = names
            notes.append(f'{shift} 按报名表「{os.path.basename(path)}[{sheet}]」'
                         f'{len(names)} 人(生成时重排考勤表名单)')
        return rosters, override, notes

    def _pick_roster(self):
        sel = self.tv_roster.selection()
        if not sel:
            self._set_status('请先在名单列表里选中一个班次')
            return
        shift = self.tv_roster.item(sel[0], 'values')[0]
        p = filedialog.askopenfilename(
            title=f'选择{shift}报名表/名单',
            filetypes=[('Excel 名单', '*.xlsx *.xlsm')])
        if not p:
            return
        try:
            names, sheet = excel_core.load_roster_file(p, shift)
        except Exception as e:
            messagebox.showerror('错误', f'读取名单失败:\n{e}')
            return
        if not names:
            messagebox.showwarning('提示', f'未从该文件中读到姓名:\n{p}')
            return
        self.roster_files[shift] = p
        self._persist_config()          # 选择即记忆为默认(下次启动沿用)
        self._refresh_roster_tree()
        self._invalidate_results()
        self._set_status(f'{shift} 已导入名单「{sheet}」{len(names)} 人, 已记忆为默认')

    def _use_template_roster(self):
        sel = self.tv_roster.selection()
        if not sel:
            self._set_status('请先在名单列表里选中一个班次')
            return
        shift = self.tv_roster.item(sel[0], 'values')[0]
        self.roster_files.pop(shift, None)
        self._persist_config()
        self._refresh_roster_tree()
        self._invalidate_results()
        self._set_status(f'{shift} 已恢复使用模板内名单')

    def _invalidate_results(self):
        """名单变化后已识别结果作废(生成时会自动补识别)。"""
        self.results.clear()
        self._refresh_photo_tree()
        self._refresh_preview()

    def _refresh_roster_tree(self):
        self.tv_roster.delete(*self.tv_roster.get_children())
        tpl = self._tpl_roster_map()
        for shift in detail_core.SHIFT_NAMES:
            path = (self.roster_files or {}).get(shift) or ''
            if path and os.path.isfile(path):
                try:
                    names, sheet = excel_core.load_roster_file(path, shift)
                except Exception as e:
                    self.tv_roster.insert('', 'end', values=(
                        shift, os.path.basename(path), '', f'读取失败: {e}'))
                    continue
                self.tv_roster.insert('', 'end', values=(
                    shift, os.path.basename(path), len(names),
                    f'工作表「{sheet}」, 生成时按此名单重排考勤表'))
            elif path:
                self.tv_roster.insert('', 'end', values=(
                    shift, os.path.basename(path), '', '文件不存在, 已改用模板内名单'))
            else:
                self.tv_roster.insert('', 'end', values=(
                    shift, '模板内名单', len(tpl.get(shift, [])), '默认'))

    def _parse_manual_text(self):
        shift = self.var_manual_shift.get()
        text = self.txt_manual.get('1.0', 'end').strip()
        if not text:
            self._set_status('请先在文本框里粘贴/输入明细(一行一条)')
            return
        rosters, _ov, _notes = self._effective_rosters()
        month_txt = self.var_month.get().strip()
        try:
            month = int(month_txt) if month_txt else None
        except ValueError:
            month = None
        entries, warns = detail_core.parse_detail_lines(
            text, tuple(rosters.get(shift, ())), month=month)
        if not entries:
            messagebox.showwarning('提示', '未解析到任何明细行, 请检查格式')
            return
        self.manual_text[shift] = text
        self.manual_results = [r for r in self.manual_results if r['shift'] != shift]
        self.manual_results.append({
            'path': f'(手工录入·{shift})', 'month': month, 'shift': shift,
            'entries': entries, 'warnings': list(warns), 'rotate': 0,
            'raw_count': len(entries)})
        self._refresh_preview()
        self._set_status(f'{shift} 手工明细 {len(entries)} 条已加入预览'
                         f'(警告 {len(warns)} 条)')

    # ---------------- 照片管理 ----------------
    def _pick_photos(self):
        ps = filedialog.askopenfilenames(
            title='选择考核明细照片',
            filetypes=[('图片', ' '.join('*' + e for e in IMAGE_EXTS))])
        self._add_photos(list(ps))

    def _on_drop(self, event):
        try:
            paths = [p for p in self.root.tk.splitlist(event.data)]
        except Exception:
            paths = [event.data]
        self._add_photos(paths)

    def _add_photos(self, paths):
        added = 0
        for p in paths:
            p = p.strip('{}')
            if not p or not os.path.isfile(p):
                continue
            if not p.lower().endswith(IMAGE_EXTS):
                continue
            if any(ph['path'] == p for ph in self.photos):
                continue
            self.photos.append({'path': p, 'override': ''})
            added += 1
        if added:
            self._refresh_photo_tree()
            self._set_status(f'已添加 {added} 张照片, 共 {len(self.photos)} 张')

    def _remove_photo(self):
        sel = self.tv_photos.selection()
        if not sel:
            return
        idx = self.tv_photos.index(sel[0])
        if 0 <= idx < len(self.photos):
            removed = self.photos.pop(idx)
            self.results.pop(removed['path'], None)
            self._refresh_photo_tree()
            self._refresh_preview()

    def _clear_photos(self):
        self.photos.clear()
        self.results.clear()
        self._refresh_photo_tree()
        self._refresh_preview()

    def _apply_override(self):
        sel = self.tv_photos.selection()
        if not sel:
            self._set_status('请先在照片列表中选中一行')
            return
        idx = self.tv_photos.index(sel[0])
        v = self.var_override.get()
        self.photos[idx]['override'] = '' if v == '自动' else v
        self.results.pop(self.photos[idx]['path'], None)  # 需重新识别
        self._refresh_photo_tree()
        self._refresh_preview()

    # ---------------- 识别(后台线程) ----------------
    def start_identify(self):
        if self.busy:
            messagebox.showinfo('提示', '正在处理中, 请稍候…')
            return
        if not self.photos:
            messagebox.showinfo('提示', '请先添加考核明细照片')
            return
        template = self.var_template.get().strip()
        if not os.path.isfile(template):
            messagebox.showerror('错误', f'模板文件不存在:\n{template}\n'
                                  '请点击「选择…」指定「煤库作业区跑分绩效汇总」xlsx 模板')
            return
        self._persist_config()
        self._set_busy(True, total=len(self.photos))
        threading.Thread(target=self._identify_worker, daemon=True).start()

    def _identify_worker(self):
        try:
            rosters, _ov, notes = self._effective_rosters()
            self.q.put(('status', '; '.join(notes) if notes else '使用模板内名单'))
            for i, ph in enumerate(self.photos):
                path, ov = ph['path'], ph['override']
                self.q.put(('status', f'正在识别({i + 1}/{len(self.photos)}): '
                                      f'{os.path.basename(path)}'))
                try:
                    res = detail_core.parse_photo(path, rosters, shift_override=ov)
                except Exception as e:
                    logger.exception('识别失败 %s', path)
                    res = {'path': path, 'month': None, 'shift': ov or None,
                           'entries': [], 'rotate': 0, 'raw_count': 0,
                           'warnings': [f'识别失败: {e}']}
                self.results[path] = res
                self.q.put(('photo_done', path))
            self.q.put(('identify_done', None))
        except Exception:
            logger.exception('识别线程异常')
            self.q.put(('fatal', traceback.format_exc()))

    # ---------------- 生成(后台线程) ----------------
    def start_generate(self):
        if self.busy:
            messagebox.showinfo('提示', '正在处理中, 请稍候…')
            return
        template = self.var_template.get().strip()
        if not os.path.isfile(template):
            messagebox.showerror('错误', f'模板文件不存在:\n{template}')
            return
        if not self.photos and not self.manual_results:
            messagebox.showinfo('提示', '请先添加考核明细照片,\n'
                                        '或在「手工录入/粘贴明细」中录入明细')
            return
        self._persist_config()
        try:
            base = int(str(self.var_base.get()).strip() or excel_core.BASE_SCORE_DEFAULT)
        except ValueError:
            messagebox.showerror('错误', '基础分必须是整数, 如 80')
            return
        pool_txt = self.var_pool.get().strip()
        try:
            pool = float(pool_txt) if pool_txt else None
        except ValueError:
            messagebox.showerror('错误', '绩效总额必须是数字, 如 10800')
            return
        month_txt = self.var_month.get().strip()
        try:
            month = int(month_txt) if month_txt else None
        except ValueError:
            messagebox.showerror('错误', '月份必须是整数, 如 7')
            return
        missing = [ph['path'] for ph in self.photos if ph['path'] not in self.results]
        rosters, override, notes = self._effective_rosters()
        if notes:
            self._show_warnings(notes + ['(以上为本次生效的名单来源)'])
        self._set_busy(True, total=len(self.photos) + 1)
        threading.Thread(target=self._generate_worker,
                         args=(template, missing, base, pool, month, rosters, override),
                         daemon=True).start()

    def _generate_worker(self, template, missing, base, pool, month,
                         rosters=None, roster_override=None):
        try:
            if rosters is None:                     # 兼容直接调用(测试)的情况
                rosters = excel_core.load_rosters(template)
            for i, path in enumerate(missing):
                ph = next(p for p in self.photos if p['path'] == path)
                self.q.put(('status', f'正在识别未识别照片({i + 1}/{len(missing)})…'))
                res = detail_core.parse_photo(path, rosters,
                                              shift_override=ph['override'])
                self.results[path] = res
                self.q.put(('photo_done', path))
            self.q.put(('status', '正在生成 Excel…'))
            all_results = [self.results[ph['path']] for ph in self.photos] + \
                list(self.manual_results)
            merged = detail_core.merge_results(all_results, default_month=month)
            m = month or merged['month']
            if not m:
                self.q.put(('fatal', '未能确定月份: 明细标题未识别到「X月」'
                                    '且未在界面填写月份, 请手动填写后重试'))
                return
            out_dir = self.var_out.get().strip() or HERE
            os.makedirs(out_dir, exist_ok=True)
            out_path = excel_core.build_output_path(out_dir, m)
            report = excel_core.fill_from_details(
                template, out_path, merged['shift_entries'], month=m,
                base_score=base, pool_amount=pool,
                roster_override=roster_override)
            self.q.put(('generate_done', (out_path, m, merged, report)))
        except Exception:
            logger.exception('生成失败')
            self.q.put(('fatal', traceback.format_exc()))

    # ---------------- 队列轮询/界面刷新 ----------------
    def _poll(self):
        try:
            while True:
                msg = self.q.get_nowait()
                kind = msg[0]
                if kind == 'status':
                    self._set_status(msg[1])
                elif kind == 'photo_done':
                    self.progress.step(1)
                    self._refresh_photo_tree()
                    self._refresh_preview()
                elif kind == 'identify_done':
                    self._set_busy(False)
                    self._set_status('识别完成, 请核对预览结果后点击「一键生成」')
                elif kind == 'generate_done':
                    self._set_busy(False)
                    out_path, m, merged, report = msg[1]
                    self.out_path = out_path
                    self.btn_open_file.config(state='normal')
                    self.btn_open_dir.config(state='normal')
                    warns = merged['warnings'] + \
                        [f'[{s}] {w}' for s, r in report.items() for w in r['warnings']]
                    self._show_warnings(warns)
                    self._set_status(f'已生成: {out_path}')
                    if warns:
                        messagebox.showinfo('生成完成',
                                            f'已生成:\n{out_path}\n\n'
                                            f'警告 {len(warns)} 条, 详见预览区下方')
                    else:
                        messagebox.showinfo('生成完成', f'已生成:\n{out_path}')
                elif kind == 'fatal':
                    self._set_busy(False)
                    self._set_status('处理失败, 详情见 paofen.log')
                    messagebox.showerror('错误', str(msg[1]))
        except queue.Empty:
            pass
        self.root.after(120, self._poll)

    def _refresh_photo_tree(self):
        self.tv_photos.delete(*self.tv_photos.get_children())
        for i, ph in enumerate(self.photos):
            res = self.results.get(ph['path'])
            name = os.path.basename(ph['path'])
            count = ''
            if res:
                shift = ph['override'] or res['shift'] or '未知'
                count = str(len(res['entries']))
                status = (f"{res['month'] or '?'}月, {len(res['entries'])} 条明细"
                          if res['entries'] else '未解析到条目')
                if res['warnings']:
                    status += f'(警告{len(res["warnings"])})'
            else:
                shift = ph['override'] or '自动'
                status = '未识别'
            self.tv_photos.insert('', 'end', values=(name, shift, count, status))

    def _refresh_preview(self):
        self.tv_entries.delete(*self.tv_entries.get_children())
        all_warns = []
        for i, ph in enumerate(self.photos):
            res = self.results.get(ph['path'])
            if not res:
                continue
            shift = ph['override'] or res['shift'] or '未知'
            for seq, e in enumerate(res['entries'], 1):
                self.tv_entries.insert('', 'end', values=(
                    f'#{i + 1}', shift, seq, f"{e['month'] or res['month']}.{e['day']}",
                    e['name_raw'], e['name'] or '未匹配', e['reason'],
                    f"{e['delta']:+g}"))
            for w in res['warnings']:
                all_warns.append(f"[{os.path.basename(ph['path'])}] {w}")
        for res in self.manual_results:          # 手工录入/粘贴的明细
            shift = res['shift'] or '未知'
            for seq, e in enumerate(res['entries'], 1):
                self.tv_entries.insert('', 'end', values=(
                    res['path'], shift, seq, f"{e['month'] or res['month']}.{e['day']}",
                    e['name_raw'], e['name'] or '未匹配', e['reason'],
                    f"{e['delta']:+g}"))
            for w in res['warnings']:
                all_warns.append(f"[{res['path']}] {w}")
        if not self.photos and not self.manual_results:
            all_warns.extend(self._effective_rosters()[2])
        self._show_warnings(all_warns)

    def _show_warnings(self, warns):
        self.txt_warn.config(state='normal')
        self.txt_warn.delete('1.0', 'end')
        if warns:
            self.txt_warn.insert('1.0', '\n'.join(warns))
        else:
            self.txt_warn.insert('1.0', '无警告')
        self.txt_warn.config(state='disabled')

    def _set_busy(self, busy, total=None):
        self.busy = busy
        state = 'disabled' if busy else 'normal'
        self.btn_generate.config(state=state)
        if total is not None:
            self.progress.config(maximum=total, value=0)
        if busy:
            self.btn_open_file.config(state='disabled')
            self.btn_open_dir.config(state='disabled')

    def _set_status(self, text):
        self.var_status.set(text)
        logger.info('%s', text)

    def _open_file(self):
        if self.out_path and os.path.isfile(self.out_path):
            os.startfile(self.out_path)

    def _open_folder(self):
        if self.out_path:
            folder = os.path.dirname(self.out_path) or HERE
            if os.path.isdir(folder):
                os.startfile(folder)


def _selftest():
    """--selftest: 环境/打包自检(依赖导入、OCR 模型、模板定位、明细解析)。

    结果同时打印到控制台并写入 exe 同目录的 selftest_result.txt, 便于打包后核对。
    """
    lines = []
    ok = True
    lines.append(f'python: {sys.version.split()[0]}  frozen={getattr(sys, "frozen", False)}')
    lines.append(f'程序目录: {HERE}')
    for mod in ('numpy', 'cv2', 'openpyxl', 'onnxruntime', 'rapidocr_onnxruntime'):
        try:
            m = __import__(mod)
            lines.append(f'[OK] import {mod} {getattr(m, "__version__", "")}')
        except Exception as e:                                # noqa: BLE001
            ok = False
            lines.append(f'[X] import {mod}: {e}')
    try:
        import tkinterdnd2
        lines.append(f'[OK] import tkinterdnd2 (支持拖拽)')
    except Exception as e:                                    # noqa: BLE001
        lines.append(f'[warn] tkinterdnd2 不可用(退化为无拖拽窗口): {e}')
    try:
        import numpy as _np
        engine = detail_core.get_engine()
        res, _ = engine(_np.full((64, 256, 3), 255, dtype=_np.uint8))
        lines.append(f'[OK] OCR 引擎可用(空图返回 {len(res or [])} 个文本框)')
    except Exception as e:                                    # noqa: BLE001
        ok = False
        lines.append(f'[X] OCR 引擎初始化失败: {e}')
    try:                                    # ORT 用的 VC 运行库来自哪里(老系统关键)
        base = getattr(sys, '_MEIPASS', HERE)
        capi = os.path.join(base, 'onnxruntime', 'capi')
        sys32 = os.path.join(os.environ.get('SystemRoot', r'C:\Windows'), 'System32')
        for dll in ('MSVCP140.dll', 'VCRUNTIME140.dll', 'VCRUNTIME140_1.dll'):
            local = os.path.join(capi, dll)
            if os.path.isfile(local):
                lines.append(f'[OK] OCR 用包内 {dll}')
            else:
                lines.append(f'[warn] 包内没有 {dll}, 将用目标机 '
                             f'{os.path.join(sys32, dll)}'
                             f'({"存在" if os.path.isfile(os.path.join(sys32, dll)) else "缺失"})'
                             ' —— 太旧会导致「找不到指定的程序」')
    except Exception as e:                                    # noqa: BLE001
        lines.append(f'[warn] 运行库检查跳过: {e}')
    try:
        t = excel_core.find_template(HERE)
        lines.append(f'[{"OK" if t else "warn"}] 模板: {t or "未找到, 可在界面点「选择…」指定"}')
    except Exception as e:                                    # noqa: BLE001
        ok = False
        lines.append(f'[X] 模板定位失败: {e}')
    d = default_out_dir()
    lines.append(f'[{"OK" if os.path.isdir(d) else "warn"}] 默认输出目录(桌面): {d}')
    if t:                                     # 完整链路自检: 明细 -> 填表 -> 读回校验
        tmp = None
        try:
            import tempfile

            import openpyxl
            roster = tuple(excel_core.load_rosters(t).get('甲班', ()))
            e2, _w = detail_core.parse_detail_lines(
                '7.2叶么菊发现现场安全隐患+3', roster)
            tmp = os.path.join(tempfile.gettempdir(), '_paofen_selftest.xlsx')
            excel_core.fill_from_details(t, tmp, {'甲班': e2},
                                         write_detail_sheet=False)
            wb = openpyxl.load_workbook(tmp)
            try:
                ws = wb['甲班']
                row = next(r for r in range(excel_core.DATA_FIRST_ROW, 60)
                           if str(ws.cell(row=r, column=2).value or '').strip() == '叶么菊')
                got = ws.cell(row=row, column=4).value          # 2日 -> D列
                formula = ws.cell(row=row, column=34).value
            finally:
                wb.close()
            good = got == 3 and formula == f'=80+SUM(C{row}:AG{row})'
            ok = ok and good
            lines.append(f'[{"OK" if good else "X"}] 填表自检: 叶么菊 2日格={got}, '
                         f'合计公式={formula}')
        except Exception as e:                                # noqa: BLE001
            ok = False
            lines.append(f'[X] 填表自检失败: {e}')
        finally:
            if tmp and os.path.isfile(tmp):
                try:
                    os.remove(tmp)
                except OSError:
                    pass
    entries, warns = detail_core.parse_detail_lines(
        '1、7.2叶么菊发现现场安全隐患+3', ('叶么菊',))
    good = len(entries) == 1 and entries[0]['name'] == '叶么菊' and entries[0]['delta'] == 3
    ok = ok and good
    lines.append(f'[{"OK" if good else "X"}] 明细解析自检: '
                 f'{entries[0] if entries else warns}')
    report = '\n'.join(lines) + f'\n结果: {"通过" if ok else "失败"}'
    print(report)
    try:
        with open(os.path.join(HERE, 'selftest_result.txt'), 'w',
                  encoding='utf-8') as f:
            f.write(report + '\n')
    except OSError as e:
        logger.warning('自检结果写入失败: %s', e)
    return 0 if ok else 1


def _ocr_diag(path):
    """--ocr <照片>: 单张照片识别诊断(在对方电脑上排障用)。

    GUI 程序没有控制台, 所以分阶段记录到 exe 同目录的 ocr_diag.txt:
      1) 环境(版本/系统/目录); 2) 读图; 3) OCR 引擎初始化;
      4) 模型推理(文本框数量与前几条文本); 5) 完整解析结果 + 异常回溯。
    哪一步出现 [X]/[!] 就是问题所在(例如引擎初始化失败 = 目标机缺运行库)。
    """
    lines = []
    ok = True

    def log(s):
        lines.append(s)

    log(f'照片: {path}')
    log(f'python: {sys.version.split()[0]}  frozen={getattr(sys, "frozen", False)}')
    log(f'程序目录: {HERE}')
    try:
        import platform
        log(f'系统: {platform.platform()} (ver {platform.version()})')
    except Exception:                                         # noqa: BLE001
        pass
    for mod in ('numpy', 'cv2', 'onnxruntime', 'rapidocr_onnxruntime'):
        try:
            m = __import__(mod)
            log(f'[OK] import {mod} {getattr(m, "__version__", "")}')
        except Exception as e:                                # noqa: BLE001
            ok = False
            log(f'[X] import {mod}: {e!r}')
    img = None
    try:
        t0 = time.time()
        img = detail_core.imread_cn(path)
        log(f'[OK] 读图: shape={img.shape} 耗时={time.time() - t0:.2f}s')
    except Exception as e:                                    # noqa: BLE001
        ok = False
        log(f'[X] 读图失败: {e!r}')
    if img is not None:
        try:
            t0 = time.time()
            engine = detail_core.get_engine()
            log(f'[OK] OCR 引擎初始化 耗时={time.time() - t0:.2f}s')
        except Exception:                                     # noqa: BLE001
            ok = False
            log('[X] OCR 引擎初始化失败:')
            log(traceback.format_exc())
        else:
            try:
                t0 = time.time()
                res, _elapse = engine(img)
                res = res or []
                log(f'[{"OK" if res else "X"}] OCR 推理: {len(res)} 个文本框 '
                    f'耗时={time.time() - t0:.2f}s')
                for item in res[:10]:
                    score = item[2] if len(item) > 2 else ''
                    log(f'      {item[0]!r} {score}')
            except Exception:                                 # noqa: BLE001
                ok = False
                log('[X] OCR 推理异常:')
                log(traceback.format_exc())
    try:
        t0 = time.time()
        tpl = excel_core.find_template(HERE)
        rosters = excel_core.load_rosters(tpl) if tpl else {}
        r = detail_core.parse_photo(path, rosters)
        log(f'[{"OK" if r["entries"] else "X"}] 解析: 月份={r["month"]} 班次={r["shift"]} '
            f'条目={len(r["entries"])} 旋转={r["rotate"]} 文本框={r["raw_count"]} '
            f'耗时={time.time() - t0:.2f}s')
        for i, e in enumerate(r['entries'], 1):
            log(f'      {i}. {e}')
        for w in r['warnings']:
            log(f'   [!] {w}')
        ok = ok and bool(r['entries'])
    except Exception:                                         # noqa: BLE001
        ok = False
        log('[X] parse_photo 异常:')
        log(traceback.format_exc())
    report = '\n'.join(lines) + f'\n结果: {"正常" if ok else "异常(见上面 [X]/[!] 行)"}'
    print(report)
    out = os.path.join(HERE, 'ocr_diag.txt')
    try:
        with open(out, 'w', encoding='utf-8') as f:
            f.write(report + '\n')
        print('诊断结果已写入: %s' % out)
    except OSError as e:
        logger.warning('诊断结果写入失败: %s', e)
    return 0 if ok else 1


def _probe_deps():
    """--probe: 逐个导入函数体检(在对方电脑上跑, 定位「找不到指定的程序」到底缺谁)。

    结果写 exe 同目录 deps_probe.txt: 包内每个 dll/pyd 引用的函数, 在这台机器上是否
    真的存在(按加载器顺序: 先看同目录那份, 再按名字找系统); 缺的几条就是根因。
    """
    import check_exe_compat
    base = getattr(sys, '_MEIPASS', HERE)
    files = []
    for root, _dirs, names in os.walk(base):
        for n in names:
            if n.lower().endswith(('.pyd', '.dll', '.exe')):
                files.append(os.path.join(root, n))
    lines = [f'python: {sys.version.split()[0]}  frozen={getattr(sys, "frozen", False)}',
             f'解包目录: {base} ({len(files)} 个二进制)']
    capi = os.path.join(base, 'onnxruntime', 'capi')
    for dll in ('MSVCP140.dll', 'VCRUNTIME140.dll', 'VCRUNTIME140_1.dll'):
        path = os.path.join(capi, dll)
        if os.path.isfile(path):
            lines.append(f'{dll}: 包内 {path} (版本 {check_exe_compat.file_version(path)})')
        else:
            lines.append(f'{dll}: 包内没有, 会用系统那份')
    missing = check_exe_compat.probe_missing_exports(
        files, extra_dirs=[os.path.join(base, d) for d in os.listdir(base)
                           if d.endswith('.libs')])
    if missing:
        lines.append(f'缺失的导入 {len(missing)} 条:')
        for path, dll, func, resolved in missing:
            lines.append('   [X] %s 需要 %s!%s (实际加载: %s)' % (
                os.path.relpath(path, base), dll, func, resolved))
    else:
        lines.append('[OK] 包内二进制引用的函数在这台机器上全部存在')
    report = '\n'.join(lines)
    print(report)
    out = os.path.join(HERE, 'deps_probe.txt')
    try:
        with open(out, 'w', encoding='utf-8') as f:
            f.write(report + '\n')
        print('体检结果已写入: %s' % out)
    except OSError as e:
        logger.warning('体检结果写入失败: %s', e)
    return 1 if missing else 0


def main():
    if '--selftest' in sys.argv:
        return _selftest()
    if '--probe' in sys.argv:                # 导入函数体检(定位「找不到指定的程序」)
        return _probe_deps()
    if '--ocr' in sys.argv:                  # 单张照片识别诊断(排障用)
        rest = [a for a in sys.argv[sys.argv.index('--ocr') + 1:]
                if not a.startswith('-')]
        if not rest:
            print('用法: 跑分绩效汇总自动生成工具.exe --ocr <照片路径>')
            return 2
        return _ocr_diag(rest[0])
    root = TkinterDnD.Tk() if TkinterDnD else tk.Tk()
    PaofenApp(root)
    root.mainloop()
    return 0


if __name__ == '__main__':
    sys.exit(main())

