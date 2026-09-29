# -*- coding: utf-8 -*-
"""打印体《考核明细》端到端验证: 渲染打印体样例图 -> OCR -> 填表 -> 逐格校验。

覆盖两种常见打印版式(均为打印体, 非手写):
  merged  : 序号 | 考核内容(整条「7.2叶么菊发现现场安全隐患+3」在同一格内)
  columns : 序号 | 日期 | 姓名 | 事由 | 分值  (第 10~12 条为 2 位序号, 覆盖「序号粘连日期」回归)

运行: python _test_printed.py   (渲染样例图需 Pillow, 仅本测试使用)
"""
import os
import sys

import openpyxl
from PIL import Image, ImageDraw, ImageFont

import detail_core
import excel_core

BASE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE, '_out')
FONT_PATH = r'C:\Windows\Fonts\simsun.ttc'   # 宋体, 模拟打印体
TEMPLATE = os.path.join(BASE, '参考文件', '煤库作业区跑分绩效汇总X月（表样）.xlsx')

# (日期, 姓名, 事由, 分值)
ROWS = (
    ('7.2', '叶么菊', '发现现场安全隐患', '+3'),
    ('7.3', '李新明', '疏通下料斗', '+3'),
    ('7.6', '靖豪', '安排零时性工作', '+5'),
    ('7.7', '王雪琴', '接班发现卫生未做', '+3'),
    ('7.10', '李新明', '两穿一带不规范', '-5'),
    ('7.11', '邓旭东', '更换岗位拉绳', '+3'),
    ('7.14', '徐天鹏', '疏通下料斗', '+3'),
    ('7.15', '曹启太', '发现岗位清扫器故障', '+3'),
    ('7.18', '黄双望', '发现设备异常', '+3'),
    ('7.19', '向双梅', '更换拉绳', '+3'),
    ('7.22', '田忠山', '两穿一带不规范', '-3'),
    ('7.23', '田本宏', '未按要求点检', '-3'),
)

# 渲染参数(接近真实手机翻拍的分辨率)
WIDTH, ROW_H, TITLE_H, FONT_SIZE, TITLE_SIZE = 1600, 74, 110, 42, 50


def expect_scores():
    """{姓名: [(日, 分值), ...]}。"""
    exp = {}
    for day, name, _reason, score in ROWS:
        exp.setdefault(name, []).append((int(day.split('.')[1]), int(score)))
    for v in exp.values():
        v.sort()
    return exp


def render(path, layout):
    """按版式渲染一张打印体样例图(path 落盘, 返回路径)。"""
    f_body = ImageFont.truetype(FONT_PATH, FONT_SIZE, index=0)
    f_title = ImageFont.truetype(FONT_PATH, TITLE_SIZE, index=0)
    if layout == 'merged':
        cols = (('序号', 140), ('考核内容', WIDTH - 80 - 140))
    else:
        cols = (('序号', 110), ('日期', 150), ('姓名', 190),
                ('事由', WIDTH - 80 - 110 - 150 - 190 - 150), ('分值', 150))
    img = Image.new('RGB', (WIDTH, TITLE_H + ROW_H * (len(ROWS) + 1) + 12), 'white')
    d = ImageDraw.Draw(img)
    d.text((WIDTH // 2 - 260, 18), '7月甲班考核明细', font=f_title, fill='black')
    xs = [40]
    for _t, cw in cols:
        xs.append(xs[-1] + cw)
    y0 = TITLE_H
    for i, (title, _cw) in enumerate(cols):
        d.rectangle([xs[i], y0, xs[i + 1], y0 + ROW_H], outline='black', width=2)
        d.text((xs[i] + 12, y0 + 14), title, font=f_body, fill='black')
    for r, (day, name, reason, score) in enumerate(ROWS, 1):
        y = y0 + r * ROW_H
        cells = ([str(r), f'{day}{name}{reason}{score}'] if layout == 'merged'
                 else [str(r), day, name, reason, score])
        for i, (_title, _cw) in enumerate(cols):
            d.rectangle([xs[i], y, xs[i + 1], y + ROW_H], outline='black', width=2)
            d.text((xs[i] + 12, y + 14), cells[i], font=f_body, fill='black')
    img.save(path)
    return path


def check_layout(layout, rosters, template):
    """渲染 -> 识别 -> 填表 -> 逐格校验, 返回是否通过。"""
    ok = True
    expect = expect_scores()
    img = render(os.path.join(OUT_DIR, f'打印体样例_{layout}.png'), layout)
    res = detail_core.parse_photo(img, rosters)
    print(f'\n===== 版式 {layout} =====')
    print(f"标题: {res['month']}月 {res['shift']} "
          f"(rotate={res['rotate']}, OCR框={res['raw_count']})")
    got = {}
    for e in res['entries']:
        got.setdefault(e['name'] or e['name_raw'], []).append((e['day'], e['delta']))
    for v in got.values():
        v.sort()
    for name in sorted(set(expect) | set(got)):
        match = expect.get(name) == got.get(name)
        ok = ok and match
        print(f"  {'OK ' if match else '[X]'} {name}: "
              f"期望{expect.get(name)} 实际{got.get(name)}")
    if res['month'] != 7 or res['shift'] != '甲班':
        ok = False
        print('  [X] 标题识别异常')
    for w in res['warnings']:
        ok = False
        print('  [X] 意外警告:', w)

    merged = detail_core.merge_results([res])
    out_path = excel_core.build_output_path(OUT_DIR, merged['month'])
    report = excel_core.fill_from_details(template, out_path,
                                          merged['shift_entries'], month=merged['month'])
    print(f'  生成: {os.path.basename(out_path)} 报告: {report}')
    for shift, rep in report.items():
        if rep['warnings']:
            ok = False
            print(f'  [X] 填表警告({shift}):', rep['warnings'])

    wb = openpyxl.load_workbook(out_path)
    try:
        ws = wb['甲班']
        for r in range(3, 21):
            name = str(ws.cell(row=r, column=2).value).strip()
            cells = [(d, int(ws.cell(row=r, column=2 + d).value))
                     for d in range(1, 32)
                     if isinstance(ws.cell(row=r, column=2 + d).value, (int, float))]
            exp = expect.get(name, [])
            match = cells == exp
            ok = ok and match
            print(f"  {'OK ' if match else '[X]'} {name}: 日格{cells}")
            for col, want in ((34, f'=80+SUM(C{r}:AG{r})'),
                              (35, f'=ROUND(AH{r}*AJ$21,2)')):
                if ws.cell(row=r, column=col).value != want:
                    ok = False
                    print(f'  [X] {name}: 第{col}列公式异常')
        ds = wb['甲班考核明细']
        print('  明细存档:')
        for r in range(3, 3 + len(res['entries'])):
            print('   ', ds.cell(row=r, column=1).value,
                  ds.cell(row=r, column=2).value)
    finally:
        wb.close()
    return ok


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    if not os.path.isfile(FONT_PATH):
        print('缺少打印体字库:', FONT_PATH)
        return 1
    rosters = excel_core.load_rosters(TEMPLATE)
    ok = True
    for layout in ('merged', 'columns'):
        ok = check_layout(layout, rosters, TEMPLATE) and ok
    print('\n打印体解析测试:', '全部通过' if ok else '存在不一致!')
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
