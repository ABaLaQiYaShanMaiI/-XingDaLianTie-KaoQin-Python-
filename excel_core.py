# -*- coding: utf-8 -*-
"""
excel_core.py - 《煤库作业区跑分绩效汇总》模板读取与按考核明细填写

模板约定(与「煤库作业区跑分绩效汇总X月（表样）.xlsx」一致):
  * 工作表: 甲班 / 乙班 / 丙班 / 丁班 + 各班「X班考核明细」共 8 张
  * 考勤表: B2='姓名', C2='1' ... AG2='31'(C~AG 列=1~31日);
    AH列=合计总分, AI列=绩效金额, AJ列=签字确认(手签, 程序不填);
    第3行起为名单, B列出现「备注」即结束; 备注行右侧:
    AH=SUM(AH3:AH20) 绩效总分, AI=绩效总额(模板自带, 如 10200), AJ=AI21/AH21 每分价值
  * 计分规则(与现场确认一致): 每人基础分 80, 考核明细的加减分写入当日格,
    合计总分 = 基础分 + 当月各日加减分之和(Excel 公式实时计算);
    绩效金额 = 合计总分 x 每分价值(Excel 公式), 程序不代算数值
"""
import logging
import os
import re
from copy import copy

import openpyxl

logger = logging.getLogger(__name__)

NAME_COL = 2          # B列: 姓名
DAY_COL_START = 3     # C列: 1日
MAX_DAY = 31
TOTAL_COL = 34        # AH列: 合计总分
AMOUNT_COL = 35       # AI列: 绩效金额
SIGN_COL = 36         # AJ列: 签字确认(手签, 程序不填)
DATA_FIRST_ROW = 3
NAME_HEADER = '姓名'
BASE_SCORE_DEFAULT = 80

# 外部「报名表/花名册」读取: 姓名单元格 = 2~4 个纯汉字, 且不是这些表头词
NAME_CELL_RE = re.compile(r'^[\u4e00-\u9fa5]{2,4}$')
ROSTER_HEADER_WORDS = (
    '姓名', '名字', '名单', '序号', '备注', '合计', '总计', '小计', '岗位', '班组',
    '班次', '工种', '工号', '电话', '手机', '身份证', '性别', '人员', '花名', '报名',
    '签字', '日期', '月份', '考勤', '工资', '绩效', '出勤', '工龄', '入职', '职位',
    '部门', '单位', '人数', '编号', '职工', '员工', '乙方', '甲方',
)
# 花名册里可能出现的「姓名」列表头
ROSTER_NAME_HEADERS = ('姓名', '名字', '员工姓名', '人员姓名', '职工姓名', '花名册')
# 优先当成「名单」的工作表名关键字
ROSTER_SHEET_KEYS = ('名单', '花名', '报名', '人员', '考勤', '职工', '员工')


def load_rosters(template_path):
    """读取模板中所有班次表的名单。返回 {工作表名: [姓名, ...]}。"""
    wb = openpyxl.load_workbook(template_path, data_only=False)
    rosters = {}
    try:
        for ws in wb.worksheets:
            b2 = str(ws.cell(row=2, column=NAME_COL).value or '').strip()
            c2 = str(ws.cell(row=2, column=DAY_COL_START).value or '').strip()
            if b2 != NAME_HEADER or c2 != '1':
                continue
            names = []
            for r in range(DATA_FIRST_ROW, 60):
                v = ws.cell(row=r, column=NAME_COL).value
                if v is None:
                    continue
                s = str(v).strip()
                if not s:
                    continue
                if s.startswith('备注'):
                    break
                names.append(s)
            if names:
                rosters[ws.title] = names
    finally:
        wb.close()
    return rosters


def find_template(start_dir=None):
    """按 程序目录 -> 上级目录 -> 参考文件子目录 找「跑分绩效汇总」xlsx 模板。"""
    here = os.path.dirname(os.path.abspath(__file__))
    dirs = [here, os.path.dirname(here), os.path.join(here, '参考文件')]
    if start_dir:
        dirs.insert(0, start_dir)
    seen = set()
    for d in dirs:
        if not d or not os.path.isdir(d) or d in seen:
            continue
        seen.add(d)
        try:
            names = sorted(os.listdir(d))
        except OSError:
            continue
        for n in names:
            if n.lower().endswith('.xlsx') and '跑分绩效汇总' in n:
                return os.path.join(d, n)
    return None


def build_output_path(out_dir, month):
    """生成输出文件路径, 同名文件自动追加 _1/_2 防止覆盖。"""
    base = f'煤库作业区跑分绩效汇总{month}月.xlsx'
    path = os.path.join(out_dir, base)
    if not os.path.exists(path):
        return path
    stem = base[:-len('.xlsx')]
    i = 1
    while True:
        p = os.path.join(out_dir, f'{stem}_{i}.xlsx')
        if not os.path.exists(p):
            return p
        i += 1


def _fmt_num(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _replace_month(text, month):
    """把标题文本中的「8月」/「X月」替换为目标月份。"""
    t = str(text)
    nt = re.sub(r'\d{1,2}\s*月', f'{month}月', t, count=1)
    if nt == t:
        nt = t.replace('X月', f'{month}月', 1)
    return nt


def _find_sheet(wb, title):
    if title in wb.sheetnames:
        return wb[title]
    for ws in wb.worksheets:  # 兜底: 工作表名包含班次名
        if title in str(ws.title):
            return ws
    return None


def _name_rows(ws):
    """{姓名: 行号} 名单区(B列, 遇「备注」止)。"""
    rows = {}
    for r in range(DATA_FIRST_ROW, 60):
        v = ws.cell(row=r, column=NAME_COL).value
        if v is None:
            continue
        s = str(v).strip()
        if not s:
            continue
        if s.startswith('备注'):
            break
        rows.setdefault(s, r)
    return rows



def _sum_row(ws):
    """定位备注行的「绩效总分」行(其 AH 单元格为 =SUM(...))。"""
    for r in range(DATA_FIRST_ROW, 60):
        v = ws.cell(row=r, column=TOTAL_COL).value
        if isinstance(v, str) and v.strip().upper().startswith('=SUM'):
            return r
    return None


def _prefer_sheets(titles, shift=None):
    """工作表尝试顺序: 指定班次 -> 含「名单/花名/报名…」关键字 -> 其余。"""
    ordered = []
    if shift and shift in titles:
        ordered.append(shift)
    ordered += [t for t in titles
                if t not in ordered and any(k in str(t) for k in ROSTER_SHEET_KEYS)]
    ordered += [t for t in titles if t not in ordered]
    return ordered


def _column_names(ws, col, max_row):
    """取某一列里按出现顺序排列的「像姓名」的值(去重、跳过表头词)。"""
    names = []
    for r in range(1, max_row + 1):
        s = str(ws.cell(row=r, column=col).value or '').strip()
        if not NAME_CELL_RE.match(s):
            continue
        if any(w in s for w in ROSTER_HEADER_WORDS):
            continue
        if s not in names:
            names.append(s)
    return names


def _roster_from_sheet(ws):
    """从一张工作表里取姓名列表(考勤表式 -> 姓名表头列 -> 像姓名最多的列)。"""
    # a) 与考勤表同版式: B2='姓名', 第3行起名单, 遇「备注」止
    if str(ws.cell(row=2, column=NAME_COL).value or '').strip() == NAME_HEADER:
        names = []
        for r in range(DATA_FIRST_ROW, 200):
            v = ws.cell(row=r, column=NAME_COL).value
            s = str(v).strip() if v is not None else ''
            if not s:
                continue
            if s.startswith('备注'):
                break
            names.append(s)
        if names:
            return names
    max_row = min(int(ws.max_row or 1), 300)
    max_col = min(int(ws.max_column or 1), 40)
    # b) 表头里直接写着「姓名/名字」的那一列
    for r in range(1, 4):
        for col in range(1, max_col + 1):
            if str(ws.cell(row=r, column=col).value or '').strip() in ROSTER_NAME_HEADERS:
                names = _column_names(ws, col, max_row)
                if names:
                    return names
    # c) 兜底: 取「像姓名」的值最多的那一列
    best = []
    for col in range(1, max_col + 1):
        names = _column_names(ws, col, max_row)
        if len(names) > len(best):
            best = names
    return best


def load_roster_file(path, shift=None):
    """读取外部「报名表/花名册」中的姓名列表。

    兼容: a) 与考勤表同版式(B2='姓名', 第3行起, 遇「备注」止);
          b) 带「姓名/名字」表头的名单; c) 任意一列连续排列的姓名。
    shift: 优先读取的工作表名(如 '甲班'); None=按表名关键字自动挑选。
    返回 (names, sheet_title); 未读到姓名时返回 ([], None)。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        for title in _prefer_sheets(list(wb.sheetnames), shift):
            names = _roster_from_sheet(wb[title])
            if names:
                return names, title
        return [], None
    finally:
        wb.close()


def _write_roster(ws, names, base_score):
    """用报名表名单重写考勤表 B 列名单(行数可增删), 返回 (warnings, cells)。

    名单变长时在备注行前插入行(复制数据行样式), 变短时删除多余行;
    之后重排 A 列序号、各行合计/金额公式, 并修正备注行的
    `=SUM(AH3:AH末行)` 与 `AJ=AI/AH` 自引用公式。
    """
    warns, cells = [], 0
    old_sum = _sum_row(ws)
    if old_sum is None:
        return ['未定位到备注行的合计公式(=SUM…), 未替换名单'], 0
    old_last = old_sum - 1                        # 模板原最后一条名单行
    new_last = DATA_FIRST_ROW + len(names) - 1    # 新名单最后一行
    if new_last > old_last:                       # 名单变长 -> 插入行
        extra = new_last - old_last
        ws.insert_rows(old_sum, extra)
        for r in range(old_sum, old_sum + extra):
            for col in range(1, SIGN_COL + 1):
                ws.cell(row=r, column=col)._style = copy(
                    ws.cell(row=DATA_FIRST_ROW, column=col)._style)
        warns.append(f'名单由 {old_last - DATA_FIRST_ROW + 1} 人变为 {len(names)} 人, '
                     f'已在考勤表插入 {extra} 行')
    elif new_last < old_last:                     # 名单变短 -> 删除多余行
        removed = old_last - new_last
        ws.delete_rows(new_last + 1, removed)
        warns.append(f'名单由 {old_last - DATA_FIRST_ROW + 1} 人变为 {len(names)} 人, '
                     f'已删除考勤表多余 {removed} 行')
    new_sum = new_last + 1                        # 备注行新位置
    base = _fmt_num(base_score)
    for i, name in enumerate(names):
        r = DATA_FIRST_ROW + i
        for col in range(1, SIGN_COL + 1):        # 清空旧名单/日格/合计/金额
            ws.cell(row=r, column=col).value = None
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=NAME_COL, value=name)
        ws.cell(row=r, column=TOTAL_COL, value=f'={base}+SUM(C{r}:AG{r})')
        ws.cell(row=r, column=AMOUNT_COL, value=f'=ROUND(AH{r}*AJ${new_sum},2)')
        cells += 4
    ws.cell(row=new_sum, column=TOTAL_COL, value=f'=SUM(AH3:AH{new_last})')
    ws.cell(row=new_sum, column=SIGN_COL, value=f'=AI{new_sum}/AH{new_sum}')
    cells += 2
    return warns, cells


def fill_from_details(template_path, out_path, shift_entries, month=None,
                      base_score=BASE_SCORE_DEFAULT, pool_amount=None,
                      write_detail_sheet=True, roster_override=None):
    """按解析好的考核明细填写模板并另存。

    shift_entries: {班次表名(如'甲班'): [entry, ...]};
      entry 为 detail_core.parse_entry 的返回值(day/name/name_raw/reason/delta/month)
    month: 目标月份(用于标题与输出文件名); None=不改标题月份
    base_score: 每人基础分(默认 80), 合计公式 = 基础分+SUM(当月各日)
    pool_amount: 绩效总额; None=保持模板自带值, 由 Excel 公式计算每人金额
    write_detail_sheet: 是否把明细同步写入「X班考核明细」工作表存档
    roster_override: {班次表名: [姓名, ...]} —— 用报名表名单替换考勤表 B 列名单
      (人数可变: 自动增/删行并修正序号、合计公式与备注行 =SUM 范围);
      未给出的班次沿用模板内名单

    返回 {班次表名: {'warnings': [str,...], 'cells': int}}
    """
    wb = openpyxl.load_workbook(template_path)
    report = {}
    try:
        for shift, entries in (shift_entries or {}).items():
            warns, cells = [], 0
            ws = _find_sheet(wb, shift)
            if ws is None:
                report[shift] = {'warnings': [f'工作簿中不存在工作表「{shift}」, 已跳过'],
                                 'cells': 0}
                continue
            # 0) 名单: 有报名表则重写 B 列名单(人数可增删), 否则沿用模板名单
            names = (roster_override or {}).get(shift)
            if names:
                rw, rc = _write_roster(ws, names, base_score)
                warns += rw
                cells += rc
                row_by_name = {n: DATA_FIRST_ROW + i for i, n in enumerate(names)}
            else:
                row_by_name = _name_rows(ws)
            sum_row = _sum_row(ws)          # 备注行(重写名单后可能已上/下移)
            if sum_row is None:
                sum_row = DATA_FIRST_ROW + len(row_by_name)
                warns.append('未定位到备注行的合计公式(=SUM…), 金额公式按末行引用, 请核对')
            # 1) 当日格写入加减分(同人同日多次自动合并求和)
            dup_seen = {}
            for e in entries or []:
                name = e.get('name')
                r = row_by_name.get(str(name).strip()) if name else None
                day = e.get('day')
                if r is None:
                    if name:
                        warns.append(f'姓名「{e.get("name_raw") or name}」'
                                     f'不在{shift}名单, 未记分(已写入明细存档)')
                    continue
                if not isinstance(day, int) or not (1 <= day <= MAX_DAY):
                    warns.append(f'{name} 日期 {day} 无效, 未记分')
                    continue
                col = DAY_COL_START + day - 1
                old = ws.cell(row=r, column=col).value
                val = e.get('delta') or 0
                if isinstance(old, (int, float)):
                    val = old + val
                    key = (name, day)
                    dup_seen[key] = dup_seen.get(key, 0) + 1
                ws.cell(row=r, column=col,
                        value=int(val) if float(val).is_integer() else float(val))
                cells += 1
            for (name, day), n in dup_seen.items():
                warns.append(f'{name} {day}日有 {n + 1} 条明细, 已合并为单格求和')
            # 2) 全员合计总分/绩效金额走 Excel 公式(包括无明细的人=纯基础分)
            base = _fmt_num(base_score)
            for r in row_by_name.values():
                ws.cell(row=r, column=TOTAL_COL,
                        value=f'={base}+SUM(C{r}:AG{r})')
                ws.cell(row=r, column=AMOUNT_COL,
                        value=f'=ROUND(AH{r}*AJ${sum_row},2)')
                cells += 2
            # 3) 绩效总额(可选): 写入备注行的 AI 单元格
            if pool_amount is not None:
                ws.cell(row=sum_row, column=AMOUNT_COL, value=pool_amount)
            # 4) 标题月份
            if month:
                t = ws.cell(row=1, column=1).value
                if t:
                    ws.cell(row=1, column=1, value=_replace_month(t, month))
            # 5) 明细存档
            if write_detail_sheet:
                ds = _find_sheet(wb, f'{shift}考核明细')
                if ds is not None:
                    if month:
                        t = ds.cell(row=1, column=1).value
                        if t:
                            ds.cell(row=1, column=1, value=_replace_month(t, month))
                    max_rows = 34  # 模板明细表预留行数
                    for i, e in enumerate(entries or []):
                        if i >= max_rows:
                            warns.append(f'明细超过 {max_rows} 条, 存档只保留前 {max_rows} 条')
                            break
                        ds.cell(row=DATA_FIRST_ROW + i, column=1, value=i + 1)
                        ds.cell(row=DATA_FIRST_ROW + i, column=2,
                                value=detail_archive_text(e, month))
                        cells += 1
                else:
                    warns.append(f'未找到「{shift}考核明细」工作表, 跳过明细存档')
            report[shift] = {'warnings': warns, 'cells': cells}
        wb.save(out_path)
        logger.info('已生成: %s', out_path)
    finally:
        wb.close()
    return report


def detail_archive_text(entry, month=None):
    """存档明细文本, 如「7.10李新明两穿一带不规范-5」。"""
    from detail_core import detail_text
    return detail_text(entry, month)
